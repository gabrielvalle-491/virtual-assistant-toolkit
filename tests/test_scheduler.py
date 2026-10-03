from datetime import UTC, date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from va_toolkit import scheduler as sc

BA = ZoneInfo("America/Argentina/Buenos_Aires")
MONDAY = date(2026, 10, 5)


def req(rid, minutes=30, days="Mon", tz="America/Argentina/Buenos_Aires", earliest="09:00", latest="18:00"):
    return sc.request_from_row({
        "request_id": rid, "attendee": f"Person {rid}", "email": f"{rid.lower()}@example.com", "topic": "Sync",
        "duration_min": minutes, "preferred_days": days, "timezone": tz, "earliest": earliest, "latest": latest,
    })


def block(day, start, end):
    return sc.block_from_row({"date": day, "start": start, "end": end}, BA)


def test_parse_days():
    assert sc.parse_days("Mon, wed;Friday") == frozenset({0, 2, 4})
    assert sc.parse_days("") == frozenset(range(5))
    with pytest.raises(ValueError):
        sc.parse_days("Funday")


def test_meetings_do_not_overlap_and_respect_buffer():
    blocks = [block(MONDAY, "09:00", "11:00")]
    meetings, unscheduled = sc.plan([req("A"), req("B"), req("C")], blocks, buffer_minutes=15)
    assert not unscheduled
    for first, second in zip(meetings, meetings[1:], strict=False):
        assert second.start >= first.end + timedelta(minutes=15)


def test_timezone_conversion_new_york():
    # 10:00 Buenos Aires (UTC-3) == 09:00 New York (EDT, UTC-4) in October
    meetings, _ = sc.plan([req("NY", 60, tz="America/New_York", earliest="09:00", latest="17:00")],
                          [block(MONDAY, "09:00", "12:00")])
    start_ny = meetings[0].start.astimezone(ZoneInfo("America/New_York"))
    assert start_ny.hour == 9 and start_ny.minute == 0
    assert meetings[0].start.astimezone(BA).hour == 10


def test_preferred_day_is_checked_in_attendee_timezone():
    # 21:00 Monday in Buenos Aires is already Tuesday in Sydney.
    r = req("SYD", 30, days="Tue", tz="Australia/Sydney", earliest="10:00", latest="12:00")
    meetings, unscheduled = sc.plan([r], [block(MONDAY, "20:00", "23:00")])
    assert not unscheduled
    local = meetings[0].start.astimezone(ZoneInfo("Australia/Sydney"))
    assert local.weekday() == 1 and local.hour == 10


def test_unschedulable_requests_are_reported_with_reason():
    blocks = [block(MONDAY, "09:00", "10:00")]
    meetings, unscheduled = sc.plan([req("A", 60), req("B", 30), req("C", 30, days="Fri")], blocks)
    reasons = {u.request.request_id: u.reason for u in unscheduled}
    assert len(meetings) == 1
    assert "already booked" in reasons["B"]
    assert "no availability" in reasons["C"]


def test_longest_meetings_are_placed_first():
    blocks = [block(MONDAY, "09:00", "10:30")]
    meetings, unscheduled = sc.plan([req("SHORT", 30), req("LONG", 90)], blocks)
    assert [m.request.request_id for m in meetings] == ["LONG"]
    assert [u.request.request_id for u in unscheduled] == ["SHORT"]


def test_ics_output_is_valid_rfc5545_shape():
    meetings, _ = sc.plan([req("R1")], [block(MONDAY, "09:00", "10:00")])
    stamp = datetime(2026, 10, 1, tzinfo=UTC)
    ics = sc.to_ics(meetings, "me@example.com", stamp)
    assert ics.startswith("BEGIN:VCALENDAR\r\n") and ics.endswith("END:VCALENDAR\r\n")
    assert "DTSTART:20261005T120000Z" in ics and "DTEND:20261005T123000Z" in ics
    assert "UID:R1@va-toolkit" in ics
    assert all(len(line.encode()) <= 75 for line in ics.split("\r\n"))


def test_run_on_samples_writes_invites_and_agenda(tmp_path, samples):
    meetings, unscheduled = sc.run(samples / "scheduler/meeting_requests.xlsx",
                                   samples / "scheduler/availability.xlsx", tmp_path, "me@example.com")
    assert len(meetings) == 6
    assert {u.request.request_id for u in unscheduled} == {"REQ-06", "REQ-08"}
    assert len(list((tmp_path / "invites").glob("*.ics"))) == 6
    assert (tmp_path / "agenda.ics").read_text().count("BEGIN:VEVENT") == 6
    wb = load_workbook(tmp_path / "agenda.xlsx")
    assert wb["Agenda"].max_row == 7 and wb["Unscheduled"].max_row == 3


def test_ics_text_is_escaped():
    assert sc._ics_escape("Budget; Q4, review\nnotes") == "Budget\\; Q4\\, review\\nnotes"
