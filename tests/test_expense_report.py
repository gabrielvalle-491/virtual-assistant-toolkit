import pandas as pd
import pytest
from openpyxl import load_workbook

from va_toolkit import expense_report as er

HEADER = "date,employee,category,vendor,description,amount,currency,receipt_id\n"


def write_csv(tmp_path, body):
    path = tmp_path / "receipts.csv"
    path.write_text(HEADER + body, encoding="utf-8")
    return path


def checked(tmp_path, body, policy=None):
    return er.apply_policy(er.load_receipts(write_csv(tmp_path, body)), policy)


def test_meals_over_daily_cap_are_flagged_per_employee_and_day(tmp_path):
    df = checked(tmp_path, (
        "2026-09-01,Ann,Meals,A,Lunch,15,USD,R1\n"
        "2026-09-01,Ann,Meals,B,Dinner,12,USD,R2\n"   # 27 > 25 on the same day
        "2026-09-01,Bob,Meals,A,Lunch,20,USD,R3\n"    # other employee, fine
        "2026-09-02,Ann,Meals,A,Lunch,25,USD,R4\n"    # exactly the cap, fine
    ))
    assert [bool(f) for f in df["flags"]] == [True, True, False, False]
    assert "27.00" in df.loc[0, "flags"]


def test_approval_threshold_missing_receipt_and_duplicate(tmp_path):
    df = checked(tmp_path, (
        "2026-09-03,Ann,Travel,Air,Flight,650,USD,R1\n"
        "2026-09-04,Ann,Transport,Taxi,Ride,10,USD,\n"
        "2026-09-05,Bob,Office Supplies,Shop,Ink,40,USD,R3\n"
        "2026-09-05,Bob,Office Supplies,Shop,Ink,40,USD,R4\n"
    ))
    assert "needs manager approval" in df.loc[0, "flags"]
    assert df.loc[1, "flags"] == "missing receipt"
    assert df.loc[2, "flags"] == ""
    assert df.loc[3, "flags"] == "possible duplicate"


def test_currency_conversion_and_unknown_currency(tmp_path):
    df = checked(tmp_path, (
        "2026-09-01,Ann,Software,Notion,Plan,10,EUR,R1\n"
        "2026-09-01,Ann,Software,Other,Plan,10,GBP,R2\n"
    ))
    assert df.loc[0, "amount_usd"] == pytest.approx(10.8)
    assert "unsupported currency GBP" in df.loc[1, "flags"]


def test_custom_policy_cap(tmp_path):
    body = "2026-09-01,Ann,Meals,A,Lunch,15,USD,R1\n"
    assert checked(tmp_path, body, er.ExpensePolicy(meals_daily_cap_usd=10)).loc[0, "flags"] != ""


def test_build_report_filters_month(tmp_path):
    df = er.load_receipts(write_csv(tmp_path, (
        "2026-08-31,Ann,Meals,A,Lunch,10,USD,R1\n"
        "2026-09-01,Ann,Meals,A,Lunch,10,USD,R2\n"
        "2026-09-02,Ann,Transport,T,Ride,5.5,USD,R3\n"
    )))
    report = er.build_report(df, "2026-09")
    assert len(report.expenses) == 2
    assert report.total_usd == 15.5
    with pytest.raises(ValueError):
        er.build_report(df, "2025-01")


def test_workbook_has_formulas_chart_and_flags(tmp_path, samples):
    report = er.run(samples / "expenses/receipts.csv", "2026-09", tmp_path)
    wb = load_workbook(tmp_path / "expense_report_2026-09.xlsx")
    assert wb.sheetnames == ["Summary", "Expenses", "Policy flags"]
    summary, expenses = wb["Summary"], wb["Expenses"]
    assert summary["C5"].value.startswith("=SUMIF(Expenses!$C$2:$C$")
    assert expenses["I2"].value == "=ROUND(F2*H2,2)"
    assert expenses.cell(expenses.max_row, 9).value.startswith("=SUM(I2:I")
    assert len(summary._charts) == 1
    assert wb["Policy flags"].max_row - 1 == len(report.flagged) == 5


def test_sample_totals_match_python_calculation(samples):
    df = er.load_receipts(samples / "expenses/receipts.csv")
    report = er.build_report(df, "2026-09")
    sept = df[pd.to_datetime(df["date"]).dt.strftime("%Y-%m") == "2026-09"]
    expected = (sept["amount"] * sept["currency"].map(er.ExpensePolicy().fx_to_usd)).round(2).sum()
    assert report.total_usd == pytest.approx(expected)
    assert report.by_category.sum() == pytest.approx(report.total_usd)
