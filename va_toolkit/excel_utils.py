"""Small openpyxl helpers shared by every module (reading tables, styling sheets)."""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN_BORDER = Border(bottom=Side(style="thin", color="BFBFBF"))


def read_table(path: str | Path, sheet: str | None = None) -> list[dict[str, Any]]:
    """Read an Excel sheet whose first row is a header into a list of dicts.

    Header names are stripped and lower-cased; fully empty rows are skipped.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return []
        keys = [str(h).strip().lower() if h is not None else "" for h in header]
        records: list[dict[str, Any]] = []
        for row in rows:
            if row is None or all(v is None or str(v).strip() == "" for v in row):
                continue
            records.append({k: v for k, v in zip(keys, row, strict=False) if k})
        return records
    finally:
        wb.close()


def write_table(ws: Worksheet, headers: Sequence[str], rows: Iterable[Sequence[Any]]) -> int:
    """Write a header row plus data rows, style the header and freeze it.

    Returns the number of data rows written.
    """
    ws.append(list(headers))
    count = 0
    for row in rows:
        ws.append(list(row))
        count += 1
    style_header(ws)
    ws.freeze_panes = "A2"
    autofit_columns(ws)
    return count


def style_header(ws: Worksheet, row: int = 1) -> None:
    """Apply the toolkit's header style to one row."""
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def autofit_columns(ws: Worksheet, min_width: int = 8, max_width: int = 60) -> None:
    """Approximate Excel's auto-fit by measuring the longest value per column."""
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = max(len(line) for line in str(cell.value).splitlines() or [""])
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, width + 2))
