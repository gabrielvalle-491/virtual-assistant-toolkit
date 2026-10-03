from datetime import date
from email import message_from_bytes, policy

import pytest
from openpyxl import load_workbook

from va_toolkit import mail_merge as mm

TEMPLATE = "Subject: Invoice {{ invoice }} for {{ company }}\n\nHi {{ first_name }}, you owe USD {{ amount | money }} by {{ due | date }}.\n"
HEADERS = ["first_name", "last_name", "email", "company", "invoice", "amount", "due"]


@pytest.mark.parametrize("email", ["a@b.co", "first.last+tag@example.co.uk", "x_y@sub.domain.org"])
def test_valid_emails(email):
    assert mm.is_valid_email(email)


@pytest.mark.parametrize("email", ["", "plainaddress", "no@tld", "a@b.c", "two@@x.com", "sp ace@x.com", None])
def test_invalid_emails(email):
    assert not mm.is_valid_email(email)


def test_parse_template_extracts_subject_and_fields():
    tpl = mm.parse_template(TEMPLATE)
    assert tpl.subject == "Invoice {{ invoice }} for {{ company }}"
    assert tpl.body.startswith("Hi {{ first_name }}")
    assert tpl.fields == {"invoice", "company", "first_name", "amount", "due"}


def test_parse_template_requires_subject():
    with pytest.raises(ValueError):
        mm.parse_template("Hello {{ name }}")


def test_validate_contact_reports_missing_fields_and_duplicates():
    required = {"first_name", "amount"}
    seen = {"dup@example.com"}
    assert mm.validate_contact({"email": "ok@example.com", "first_name": "A", "amount": 1}, required, seen) == []
    issues = mm.validate_contact({"email": "DUP@example.com", "first_name": " ", "amount": None}, required, seen)
    assert issues == ["duplicate email", "missing field(s): amount, first_name"]


def test_merge_writes_drafts_and_skips_bad_rows(tmp_path, make_xlsx):
    contacts = make_xlsx("c.xlsx", HEADERS, [
        ["Ana", "Paz", "ana@example.com", "Acme", "INV-1", 1234.5, date(2026, 10, 15)],
        ["Bo", "Lee", "bad-email", "Beta", "INV-2", 10, date(2026, 10, 15)],
        ["", "X", "x@example.com", "Gamma", "INV-3", 10, date(2026, 10, 15)],
        ["Ana", "Paz", "ANA@example.com", "Acme", "INV-4", 5, date(2026, 10, 15)],
    ])
    tpl = tmp_path / "t.txt"
    tpl.write_text(TEMPLATE, encoding="utf-8")
    results = mm.run(contacts, tpl, tmp_path / "out", "Me <me@example.com>")

    assert [r.status for r in results] == ["READY", "SKIPPED", "SKIPPED", "SKIPPED"]
    drafts = list((tmp_path / "out" / "drafts").glob("*.eml"))
    assert len(drafts) == 1
    msg = message_from_bytes(drafts[0].read_bytes(), policy=policy.default)
    assert msg["To"] == "ana@example.com"
    assert msg["Subject"] == "Invoice INV-1 for Acme"
    assert msg["X-Unsent"] == "1"
    assert "USD 1,234.50 by October 15, 2026" in msg.get_content()


def test_summary_workbook_lists_every_contact(tmp_path, make_xlsx):
    contacts = make_xlsx("c.xlsx", HEADERS, [
        ["Ana", "Paz", "ana@example.com", "Acme", "INV-1", 1, date(2026, 1, 1)],
        ["Bo", "Lee", "", "Beta", "INV-2", 1, date(2026, 1, 1)],
    ])
    tpl = tmp_path / "t.txt"
    tpl.write_text(TEMPLATE, encoding="utf-8")
    mm.run(contacts, tpl, tmp_path / "out", "me@example.com")
    ws = load_workbook(tmp_path / "out" / "mail_merge_summary.xlsx").active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert rows[1][3] == "SKIPPED" and rows[1][4] == "missing email"


def test_sample_data_produces_expected_split(tmp_path, samples):
    results = mm.run(samples / "mail_merge/contacts.xlsx", samples / "mail_merge/template.txt", tmp_path, "me@example.com")
    assert sum(r.status == "READY" for r in results) == 7
    assert sum(r.status == "SKIPPED" for r in results) == 5
